#!/usr/bin/env python3
"""
check_master_list.py — Cross-check of a KB base's master-index spreadsheet
("Lista Mestra").

This script is FIXED (not rewritten on every run) so the consistency check
always runs the same way, with the same recommendation criteria — instead of
depending on a freshly generated script that could vary from one audit to
the next.

WHAT IT CHECKS
--------------
1. Broken references: every master-index row whose "File" column points to a
   file that doesn't actually exist in the folder.
2. "Orphaned" files: real files in the base that don't appear in any
   master-index row (files inside a folder whose name ends with "- Files",
   e.g. "KB-017 - Autopilot - Files/", are treated as attachments of that KB
   document and are never flagged this way — see ATTACHMENT_DIR_SUFFIXES).
3. Duplicate/shared codes: two different rows using the same code (except
   the "—" marker, which can repeat freely for documents without a code).
4. Name divergence: for every .docx document, compares the name derived from
   4 sources — file name, title (1st line of the body), header, and the
   master-index spreadsheet's "Document" column — and flags when they don't
   match.
5. Version divergence: compares the header's version, the last row of the
   Version History table, and the master-index spreadsheet's "Version"
   column.
6. Version History table in the legacy format (No. | Revision Date |
   Revision | Reviewer, with no semantic-version column) — reported as its
   own item, so it doesn't produce a false version divergence by comparing a
   date against a version number.
7. Header missing the tab separator between name and version — either glued
   ("Namev1.2") or separated by a space ("Name v1.2"). Reported as its own
   item, so it doesn't show up disguised as a generic version divergence, or
   worse, as a NAME divergence whose recommendation tells you to copy the
   broken name into the other three sources. The spaced form is only treated
   as a version when the file name or the Version History agrees that is what
   it is — otherwise "Migracao Office v2" would lose its own name.
8. Header with no version at all. Section 4 requires one in every header, and
   an absent version used to be invisible: the version comparison drops empty
   sources before comparing, so {None, "1.0", "1.0"} collapsed to one value
   and reported nothing.
9. Documents the spreadsheet points at that exist but cannot be parsed as Word
   documents. Reported once, under their own heading: such a file used to
   produce three findings at once — a "broken reference" claiming a file that
   plainly exists doesn't, plus a structure and a formatting violation.

In addition, at the end of the report, an INFORMATIONAL section (never
counted as a problem) lists files whose name contains a non-ASCII character
(accents, etc.) — a portability map for a possible future migration, not a
base inconsistency.

Every file-name comparison (master-index spreadsheet vs. real files) is done
after normalizing the strings to Unicode NFC (see the nfc() function), so
that the same name saved as NFC and as NFD isn't flagged as a broken
reference just because the underlying bytes differ.

RECOMMENDATION ON NAME DIVERGENCE
----------------------------------
The script does NOT decide or apply the fix on its own. It points to a
recommendation based on which source was changed most recently:

- Compares the .docx file's modification date (file mtime, which reflects
  the last time the title/header were edited) with the date in the
  master-index spreadsheet's "Last Updated" column (parsed whether it's a
  real Excel date or plain text — see parse_date_value()).
- The most recently changed source is suggested as the "recommended" name —
  but the report always lists EVERY variant found, with its origin and date,
  so the person running it can decide which one to keep everywhere (file,
  title, header, and master-index spreadsheet).
- If the dates don't allow a decision (e.g. equal, or unavailable), the
  script marks the divergence as "no automatic recommendation — manual
  decision needed" instead of guessing. If the spreadsheet's "Last Updated"
  cell has a value but it couldn't be parsed as a date at all, the report
  says so explicitly instead of silently falling back to "unavailable".

COLUMN LOOKUP AND EXIT CODE
----------------------------
Master-index spreadsheet rows are read by matching each column's NAME (see
COLUMN_SYNONYMS), not by a fixed position — the header row itself is found
by scanning for a recognizable "Code"/"Código" column rather than assuming
it's always row 2. A sheet with no recognizable header row is skipped with a
WARNING rather than silently read as if its columns were in the usual order.

The process exit code reflects the result: 0 when no problems were found, 1
when at least one was — so this script can gate a batch edit or a CI-style
check on a clean run instead of always reporting success regardless of what
it found.

HOW TO USE
----------
    python3 check_master_list.py <base_path> [<spreadsheet_name.xlsx>]

- <base_path>: the knowledge base's root directory (contains the category
  subfolders plus the master-index spreadsheet).
- <spreadsheet_name.xlsx>: optional. If omitted, the script looks for the
  single matching master-index spreadsheet file at the base's root (see
  MASTER_LIST_NAME_PATTERNS below for the recognized naming conventions).

Requirements: python-docx, openpyxl (pip install python-docx openpyxl).

OUTPUT
------
A text report on the console, grouped by problem type, always ending with a
summary of how many items of each type were found. Nothing is changed in any
file — this script is read-only/diagnostic only.

NOTE ON LANGUAGE: this script's own docstring, comments, and printed report
are in English so the tool itself is easy to read, maintain, and share
across teams. The KB base it inspects can be in any target document
language (see the kb-compiler skill's Section 0/6) — the patterns used to
recognize structural elements (sheet names, column headers) below
deliberately match BOTH the Portuguese terms this plugin was originally
built around (Hasbro Brazil's live base) and their common English
equivalents, so the same fixed script keeps working for either without a
fork. If a base uses yet another language for these structural terms, widen
the matching sets below rather than writing a new script.
"""

import sys
import os
import glob
import json
import re
import datetime
import unicodedata

# Windows defaults stdout to the locale encoding (cp1252) whenever it is not a
# console, so redirecting or piping this report died with a UnicodeEncodeError
# on the first section title containing "→" — after printing two lines, and
# exiting 1, which reads as "problems found" rather than "crashed". The report
# is full of "→", "—" and "⚠", so the stream is pinned to UTF-8 instead.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

try:
    import docx
    from docx.oxml.ns import qn
except ImportError:
    print("ERROR: missing 'python-docx'. Install with: pip install -r requirements.txt")
    sys.exit(2)

try:
    import openpyxl
except ImportError:
    print("ERROR: missing 'openpyxl'. Install with: pip install -r requirements.txt")
    sys.exit(2)


# ---------------------------------------------------------------------------
# Language vocabulary. Every localized term lives in languages.json, NOT here.
#
# Section 10 of the standard says a base may be authored in any language, but
# the code used to hardcode English and Portuguese across eight separate
# tables — so "any language" actually meant "either of two, and only by editing
# Python". The recognition sets below are now built by UNIONING every language
# in that file, which is what lets one fixed script read a base in any of them.
# Adding a language is a data edit; nothing here changes.
# ---------------------------------------------------------------------------
LANGUAGES_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "languages.json")


def load_languages(path=LANGUAGES_FILE):
    """Returns {code: vocabulary}. A missing or broken file is fatal: guessing
    the vocabulary would mean silently auditing against the wrong terms."""
    try:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
    except (OSError, ValueError) as e:
        print("ERROR: couldn't read %s (%s). It defines every localized term "
              "the audit relies on." % (os.path.basename(path), e))
        sys.exit(2)
    langs = data.get("languages")
    if not isinstance(langs, dict) or not langs:
        print("ERROR: %s has no 'languages' section." % os.path.basename(path))
        sys.exit(2)
    # Validate the keys the module-level tables below index DIRECTLY. Adding a
    # language is documented as a data edit to this file, so a half-finished
    # entry is an expected mistake — and without this it surfaced as a bare
    # `KeyError` traceback at import time, which exits 1. That is this tool's
    # own "problems found" code, so a broken vocabulary was indistinguishable
    # from a completed audit in any wrapper keying on the exit status. The
    # __main__ handler cannot catch it: the failure happens before main runs.
    for code, vocab in sorted(langs.items()):
        if not isinstance(vocab, dict):
            print("ERROR: language %r in %s is not an object."
                  % (code, os.path.basename(path)))
            sys.exit(2)
        missing = [k for k in ("overview_sheet", "spreadsheet_name")
                   if not str(vocab.get(k) or "").strip()]
        if missing:
            print("ERROR: language %r in %s is missing required key(s): %s. "
                  "Every language needs at least these, because the audit's "
                  "sheet-skipping and spreadsheet-detection tables are built "
                  "from them."
                  % (code, os.path.basename(path), ", ".join(missing)))
            sys.exit(2)
    return langs


LANGUAGES = load_languages()


def _union(*path):
    """Collects a nested list of terms across every language, lowercased."""
    out = set()
    for vocab in LANGUAGES.values():
        node = vocab
        for key in path:
            node = node.get(key, {}) if isinstance(node, dict) else {}
        if isinstance(node, list):
            out.update(t.strip().lower() for t in node)
    return out


def _union_map(section):
    """{key: (term, ...)} merged across languages — e.g. every spelling of the
    'code' column, in every language listed."""
    merged = {}
    for vocab in LANGUAGES.values():
        for key, terms in (vocab.get(section) or {}).items():
            merged.setdefault(key, set()).update(t.strip().lower() for t in terms)
    return {k: tuple(sorted(v)) for k, v in merged.items()}


# Overview/summary tab names to skip when scanning category tabs.
#
# Compared case-insensitively. Matching the literal spelling meant a tab renamed
# "OVERVIEW" stopped being recognized as the Overview: it was then read as a
# category sheet, found no Code column, and was skipped — which silently lost
# the tab -> category map that exists to prevent false folder/tab mismatches.
SKIP_SHEETS = {v["overview_sheet"].strip().lower() for v in LANGUAGES.values()} | {
    alias.strip().lower()
    for v in LANGUAGES.values() for alias in v.get("overview_sheet_aliases", [])
}


def is_overview_sheet(sheet_name):
    """True for the Overview/summary tab in any language, in any casing."""
    return (sheet_name or "").strip().lower() in SKIP_SHEETS

# Master-index spreadsheet naming conventions to auto-detect.
MASTER_LIST_NAME_PATTERNS = tuple(
    sorted({"%s*.xlsx" % v["spreadsheet_name"] for v in LANGUAGES.values()}
           | {"Master Index*.xlsx"})
)

NO_CODE_MARKER = "—"

# The three Status values Section 7 allows, with the spellings the tooling
# recognizes. Compared lowercased but WITH accents: the spreadsheet's own
# spelling is what matters, and "Em revisao" missing its cedilla is exactly the
# typo this list exists to catch.
#
# A wrong Status is not cosmetic. It silently disables logic: the Overview's
# COUNTIF formulas stop matching, so the category totals under-report, and a
# misspelled legacy label makes a superseded document fall through to the
# structure checks it was meant to be exempt from.
STATUS_LABELS = _union_map("status")
VALID_STATUS_VALUES = {v for spellings in STATUS_LABELS.values() for v in spellings}

# Status values that mark a document as superseded. Structure and Version
# History sequence checks are skipped for these. Derived from STATUS_LABELS
# rather than repeated, so the two can't drift apart.
LEGACY_STATUS_LABELS = set(STATUS_LABELS.get("legacy", ()))

# ---------------------------------------------------------------------------
# Formatting constants — these mirror Section 4 of SKILL.md ("Exact Word
# formatting"). They live here so the spec and the checker can be compared
# side by side; if Section 4 changes, change these too.
# ---------------------------------------------------------------------------
EXPECTED_MARGIN_INCHES = 0.75
EXPECTED_TITLE_SIZE_PT = 22.0
EXPECTED_TITLE_COLOR = "1F3864"
EXPECTED_CALLOUT_SHADING = "FFF3CD"
EXPECTED_CALLOUT_TEXT_COLOR = "7B5C00"
CALLOUT_PREFIX = "⚠"
HEADING_STYLE_MARKERS = tuple(sorted(_union("heading_style_markers")))

# Sections whose body must be a list, and which kind. Checked only when the
# section actually has body content — an empty section can't be judged.
_SECTIONS = _union_map("sections")
# Prerequisites and Verification must be a LIST, but either kind: bulleted or
# numbered. The standard used to say bulleted, and a real base showed why that
# was wrong — it used numbered lists essentially everywhere, so the rule was
# describing a practice nobody followed. What actually matters is that the
# items are a list rather than a wall of prose.
ANY_LIST_SECTIONS = tuple(sorted(set(_SECTIONS.get("prerequisites", ()))
                                 | set(_SECTIONS.get("verification", ()))))
# Step by Step stays decimal: the steps are ordered, and a bullet loses that.
NUMBERED_SECTIONS = tuple(sorted(_SECTIONS.get("steps", ())))
DASHES = ("\u2013", "\u2014")  # en dash, em dash — the standard requires a plain hyphen

# A folder whose name ends with one of these (case-insensitive) is treated as
# an attachment/supporting-files folder for a KB document (e.g.
# "KB-017 - Autopilot - Files/"), not a set of independently indexable items.
# Files inside it are never flagged as "orphaned" even though they don't have
# their own master-index row — they're dependencies of the KB they sit under,
# not documents in their own right.
ATTACHMENT_DIR_SUFFIXES = ("- files",)

# The folder a superseded document lives in, per Section 1. Taken from the
# vocabulary file so adding a language stays a data edit, with the English term
# as a floor in case a language omits it.
LEGACY_FOLDER_NAMES = {t for t in _union("legacy_folders")} | {"legacy"}

# Column names recognized in the master-index spreadsheet's header row, in
# Portuguese (this plugin's original base) and English, so rows are read by
# column NAME instead of a fixed position — a column inserted or reordered in
# the spreadsheet won't silently shift every other column's data.
COLUMN_SYNONYMS = _union_map("columns")

# Version History table columns, in both the current and the legacy layouts.
HISTORY_COLUMNS = _union_map("history_columns")
LEGACY_HISTORY_COLUMNS = _union_map("legacy_history_columns")

# The Status spellings a person should actually see in an error message: the
# written forms, not the internal keys ("in_review").
STATUS_WRITTEN = sorted({w for v in LANGUAGES.values()
                         for w in (v.get("status_written") or {}).values()})


def nfc(text):
    """
    Normalize to Unicode NFC (composed form). Visually identical file names
    can be saved as NFC or NFD (common when a file passes through
    macOS/OneDrive/different OSes) — e.g. an accented character as a single
    code point (NFC) vs. a base letter + combining accent (NFD) are
    different bytes and won't match in a naive string comparison, producing
    false "broken reference" positives. Every string used as a file-name
    comparison key goes through this first.
    """
    if not text:
        return text
    return unicodedata.normalize("NFC", text)


# Per-base configuration written by bootstrap_master_list.py. It records where
# this base's master index lives (and the setup answers that go with it), so
# every later run uses the SAME spreadsheet instead of re-deriving it from a
# file-name pattern. That is what makes a custom name or a location outside the
# base folder work as a continuous insertion point.
CONFIG_FILENAME = ".kb-compiler.json"


def stdin_is_interactive():
    """
    True only when there is a human who can actually answer a prompt.

    `isatty()` alone is wrong on Windows: stdin redirected from NUL is a
    character device, so isatty() reports a TERMINAL. A scheduled task or CI
    run therefore printed a numbered menu to nobody, and only discovered the
    truth when input() raised EOFError — by which point the menu was already on
    screen, immediately followed by "running non-interactively", which reads as
    the tool contradicting itself.

    A real console answers GetConsoleMode; NUL does not, despite both being
    character devices. That is the discriminator, and it costs one call. The
    EOFError handlers downstream stay as a second line of defence.
    """
    try:
        if not sys.stdin.isatty():
            return False
    except (AttributeError, ValueError):
        return False          # detached stdin: nobody there
    if os.name != "nt":
        return True
    try:
        import ctypes
        import msvcrt
        handle = msvcrt.get_osfhandle(sys.stdin.fileno())
        mode = ctypes.c_ulong()
        return bool(ctypes.windll.kernel32.GetConsoleMode(handle, ctypes.byref(mode)))
    except Exception:
        # Couldn't tell. Assume interactive: guessing the other way would
        # suppress a prompt a real person is sitting there waiting to answer,
        # whereas guessing this way is still caught by the EOFError handlers.
        return True


def read_config(base_dir):
    """Returns the base's config dict, or {} when there is none / it's unreadable."""
    path = os.path.join(base_dir, CONFIG_FILENAME)
    if not os.path.isfile(path):
        return {}
    try:
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except (ValueError, OSError):
        return {}


def write_config(base_dir, **values):
    """Merges values into the base's config file and returns the merged dict."""
    data = read_config(base_dir)
    data.update({k: v for k, v in values.items() if v is not None})
    path = os.path.join(base_dir, CONFIG_FILENAME)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
        f.write("\n")
    return data


def resolve_configured_master_list(base_dir):
    """
    Returns the configured master-index path, or None. A configured path that
    no longer exists is an explicit error rather than a silent fallback to
    pattern matching — falling back would audit a different file than the one
    the base was set up to use.
    """
    configured = read_config(base_dir).get("master_index")
    if not configured:
        return None
    path = configured if os.path.isabs(configured) else os.path.join(base_dir, configured)
    if not os.path.isfile(path):
        raise FileNotFoundError(
            "%s points at '%s', which no longer exists. Fix the path in that file, "
            "or delete it to fall back to auto-detection." % (CONFIG_FILENAME, configured)
        )
    return path


def find_master_list(base_dir, explicit_name=None):
    # An explicitly passed name always wins; then the base's own configuration;
    # only then pattern matching.
    if not explicit_name:
        configured = resolve_configured_master_list(base_dir)
        if configured:
            return configured

    if explicit_name:
        path = os.path.join(base_dir, explicit_name)
        if not os.path.isfile(path):
            raise FileNotFoundError(f"Spreadsheet '{explicit_name}' not found in {base_dir}")
        return path

    # Accumulate matches across EVERY naming pattern before deciding. Breaking
    # on the first pattern that matched meant a base holding both a
    # "Lista Mestra" and a "Master List" silently audited only one of them —
    # a clean report you would have believed.
    candidates = []
    for pattern in MASTER_LIST_NAME_PATTERNS:
        candidates.extend(glob.glob(os.path.join(base_dir, pattern)))
    candidates = sorted(set(candidates))

    if not candidates:
        raise FileNotFoundError(
            f"No master-index spreadsheet found in {base_dir} (tried: "
            f"{', '.join(MASTER_LIST_NAME_PATTERNS)}). Pass the file name explicitly "
            "as the second argument."
        )
    if len(candidates) > 1:
        raise RuntimeError(
            f"More than one master-index spreadsheet found in {base_dir}: "
            f"{candidates}. Pass the one to use as the second argument."
        )
    return candidates[0]


def is_master_list_filename(basename):
    """
    True only for a file that could actually BE a master index.

    The extension check is the point: matching on the name prefix alone meant a
    document called "Master List Guidelines.docx" was treated as the index and
    silently excluded from orphan detection — an unindexed document that never
    appeared in the report at all. Compared case-insensitively, since the
    patterns are globbed case-insensitively on Windows but not on Linux.
    """
    lowered = (basename or "").lower()
    if not lowered.endswith(".xlsx"):
        return False
    for pattern in MASTER_LIST_NAME_PATTERNS:
        prefix = pattern.split("*")[0].lower()
        if lowered.startswith(prefix):
            return True
    return False


# Files that are never indexable documents: Office lock files, and the junk
# Windows/macOS leave in a synced folder. Without this, auditing a base while a
# single document was open in Word reported its "~$" lock file as an orphan —
# a finding that disappears on its own and makes the report look unreliable.
IGNORED_BASENAMES = frozenset(("thumbs.db", ".ds_store", "desktop.ini"))


def is_ignorable_file(basename):
    lowered = (basename or "").strip().lower()
    if lowered in IGNORED_BASENAMES:
        return True
    # "~$Doc.docx" (Word/Excel) and ".~lock.Doc.docx#" (LibreOffice).
    return lowered.startswith("~$") or lowered.startswith(".~lock.")


def is_inside_attachment_folder(rel_path):
    """
    True if any directory component of rel_path (everything but the file's
    own name) ends with one of ATTACHMENT_DIR_SUFFIXES — e.g. a file under
    "KB-017 - Autopilot - Files/GetAutoPilot/tool.zip" is an attachment of
    KB-017, not a standalone item that belongs in the master-index
    spreadsheet on its own.
    """
    parts = rel_path.split("/")[:-1]
    for part in parts:
        if part.strip().lower().endswith(ATTACHMENT_DIR_SUFFIXES):
            return True
    return False


def find_header_row(rows):
    """
    Scans a sheet's rows (as returned by iter_rows(values_only=True)) for the
    one that is the column-header row, identified by having a cell that
    matches the "code" column's synonyms — rather than assuming it's always
    row 2. Returns (row_index, lowered_cell_values) or (None, None) if no
    such row is found.
    """
    for i, row in enumerate(rows):
        if not row:
            continue
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        if any(c in COLUMN_SYNONYMS["code"] for c in cells):
            return i, cells
    return None, None


def map_columns(header_cells):
    """Maps each COLUMN_SYNONYMS key to the column index found in header_cells."""
    mapping = {}
    for key, synonyms in COLUMN_SYNONYMS.items():
        for i, c in enumerate(header_cells):
            if c in synonyms:
                mapping[key] = i
                break
    return mapping


_SLASH_DATE = re.compile(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})$")


def date_is_ambiguous(value):
    """
    True for a TEXT date whose first two fields could each be the day or the
    month, and which therefore means two different dates depending on the
    convention — "03/04/2026" is 3 April or 4 March.

    A real Excel date cell is never ambiguous: it carries a date, not a
    rendering. Neither is "13/04/2026", where only one reading is a valid
    month, nor "05/05/2026", where both readings agree.
    """
    if value is None or isinstance(value, (datetime.date, datetime.datetime)):
        return False
    m = _SLASH_DATE.match(str(value).strip())
    if not m:
        return False
    first, second = int(m.group(1)), int(m.group(2))
    return first != second and first <= 12 and second <= 12


def parse_date_value(value, day_first=True):
    """
    Best-effort parse of a spreadsheet date-like cell into a date, whether
    it's a real datetime/date (the normal case for an Excel date cell) or
    plain text (e.g. a "Last Updated" column typed in as a string instead of
    a real date). Returns None if the value is empty or couldn't be parsed
    as a date in any of the tried formats.

    `day_first` picks which reading wins for an ambiguous text date. The
    caller passes the base's own convention, taken from the language recorded
    in its configuration: trying %d/%m/%Y unconditionally meant an English
    base's "03/04/2026" was read as 3 April. Where it genuinely cannot be
    settled, date_is_ambiguous() lets the report say so rather than presenting
    one reading as fact.
    """
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value).strip()
    if not text:
        return None
    orders = (("%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y") if day_first
              else ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y"))
    for fmt in orders:
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def index_real_files(base_dir):
    """
    Returns (by_basename, by_relative_path) for every real file. Keys are
    normalized to NFC (see nfc()) so the comparison against the
    master-index spreadsheet's names doesn't depend on which Unicode form
    either side happens to use.
    """
    by_basename = {}
    by_path = {}
    for root, _dirs, files in os.walk(base_dir):
        for f in files:
            if is_ignorable_file(f):
                continue
            full = os.path.join(root, f)
            rel = os.path.relpath(full, base_dir).replace(os.sep, "/")
            f_norm = nfc(f)
            rel_norm = nfc(rel)
            by_basename.setdefault(f_norm, []).append(rel_norm)
            by_path[rel_norm] = full
    return by_basename, by_path


def strip_code_prefix(text):
    """Removes a 'KB-XXX - ' (or similar) prefix from the start of the text, if present."""
    if not text:
        return text
    text = text.strip()
    # Accept the dash variants too. The standard requires a plain hyphen and a
    # header using an en/em dash is reported separately — but failing to strip
    # the prefix here would ALSO surface it as a bogus name divergence, i.e.
    # the same defect counted twice, with a recommendation to propagate the
    # broken form to the other sources.
    for separator in (" - ", " \u2013 ", " \u2014 "):
        parts = text.split(separator, 1)
        if len(parts) == 2 and parts[0].upper().startswith("KB-"):
            return parts[1].strip()
    return text


def classify_table_header(header_cells):
    """
    Decides whether the Version History table's first row is a recognizable
    column header, and in which format:
      - "new": columns Author | Version | Date | Description (of the Change)
      - "legacy": columns No. | Revision Date | Revision | Reviewer
      - None: not a header row (it's already data), or an unrecognized format
    Returns (format, version_index_or_None, description_index_or_None).

    Matches both the Portuguese terms (Autor/Versão/Revisor/Revisão) this
    plugin was built around and their English equivalents
    (Author/Version/Reviewer/Revision), so the same script works for a base
    in either language.
    """
    if not header_cells or len(header_cells) < 4:
        return None, None, None
    lowered = [c.strip().lower() for c in header_cells]

    def find(terms, exact=False):
        """Index of the first cell matching any term, or None."""
        for i, c in enumerate(lowered):
            for term in terms:
                if (c == term) if exact else (term in c):
                    return i
        return None

    # new format: a column clearly called "version" AND an "author" one
    idx_version = find(HISTORY_COLUMNS.get("version", ()))
    idx_author = find(HISTORY_COLUMNS.get("author", ()))
    if idx_version is not None and idx_author is not None:
        idx_description = find(HISTORY_COLUMNS.get("description", ()))
        return "new", idx_version, idx_description

    # legacy format: "reviewer" plus "revision" as a column NAME, not as data —
    # hence the exact match, so a cell reading "Revision Date" doesn't count.
    if find(LEGACY_HISTORY_COLUMNS.get("reviewer", ())) is not None:
        idx_description = find(LEGACY_HISTORY_COLUMNS.get("revision", ()), exact=True)
        if idx_description is not None:
            return "legacy", None, idx_description
    return None, None, None


def _file_stem(full_path):
    """The document's own file name without extension, NFC-normalized."""
    return nfc(os.path.splitext(os.path.basename(full_path))[0]).strip()


def _parse_header(signals, hdr_text, full_path):
    """
    Splits the header into name and version, and classifies the separator.

    Section 4 mandates "<name><TAB>v<version>". Three broken shapes have to be
    told apart from a legitimate name that merely ends in something
    version-shaped ("... Office v2"):

      - glued:  "Namev1.0"   — no separator at all
      - spaced: "Name v1.0"  — a space where the tab belongs
      - no version in the header at all

    The spaced case is the one that needs care. Treating every trailing "v<n>"
    as a version would misread "Migracao Office v2"; treating none of them as a
    version left a space-separated header — the likeliest way to get this wrong
    by hand — reported as a NAME divergence instead, whose recommendation told
    you to propagate "Alpha v1.0" into the file name and the title. So the
    trailing token counts as a version only when another source agrees: the
    file name matches the remainder, or the Version History records exactly
    that version.
    """
    signals["header"] = hdr_text

    if "\t" in hdr_text:
        signals["header_name"] = hdr_text.split("\t")[0].strip()
        tail = hdr_text.split("\t")[-1].strip()
        if tail.lower().startswith("v") and tail[1:2].isdigit():
            signals["header_version"] = tail[1:].strip()
        else:
            signals["header_missing_version"] = True
        return

    # "Namev1.2": glued straight onto a word character. The lookbehind is what
    # keeps "Migracao Office v2" out of this branch.
    m = re.search(r"(?<=\w)v\d+(\.\d+)?$", hdr_text, re.IGNORECASE)
    if m:
        signals["header_missing_tab"] = True
        signals["header_version"] = m.group(0)[1:]
        # The regex knows where the version starts, so the name is everything
        # before it. Leaving the version glued on would report the SAME defect
        # twice — once correctly, once as a bogus name divergence.
        signals["header_name"] = hdr_text[: m.start()].strip()
        return

    # "Name v1.2": a space where the tab belongs.
    m = re.search(r"\s+v(\d+(?:\.\d+)?)$", hdr_text, re.IGNORECASE)
    if m:
        candidate_name = hdr_text[: m.start()].strip()
        candidate_version = m.group(1)
        agrees_with_file = nfc(candidate_name) == _file_stem(full_path)
        agrees_with_history = (
            signals["history_version"] is not None
            and candidate_version == signals["history_version"]
        )
        if agrees_with_file or agrees_with_history:
            signals["header_missing_tab"] = True
            signals["header_version"] = candidate_version
            signals["header_name"] = candidate_name
            return

    # Nothing version-shaped to split off: the whole header is the name, and it
    # carries no version — which Section 4 requires it to.
    signals["header_name"] = hdr_text
    signals["header_missing_version"] = True


def read_docx_signals(full_path, doc=None):
    """
    Extracts the title, header, header version, footer label and Version
    History signals. Pass `doc` to reuse an already-open Document instead of
    parsing the same file a second time.
    """
    signals = {
        "title": None,
        "header": None,
        "header_name": None,
        "header_version": None,
        "header_missing_tab": False,
        "header_missing_version": False,
        "footer_label": None,
        "history_version": None,
        "history_versions": [],
        "history_description": None,
        "table_format": None,  # "new", "legacy", None (unrecognized/empty)
        "error": None,
    }
    try:
        d = doc if doc is not None else docx.Document(full_path)
        # First NON-EMPTY paragraph: a blank line at the top of the body used
        # to make the title read as "", which silently dropped it out of the
        # four-source name comparison (turning it into a three-source one).
        title_para = next((p for p in d.paragraphs if p.text.strip()), None)
        if title_para is not None:
            signals["title"] = title_para.text.strip()

        # The Version History is read BEFORE the header, because deciding
        # whether a trailing "v1.0" in the header is a glued-on version or part
        # of the document's own name needs the version the document claims.
        if d.tables:
            last_table = d.tables[-1]
            if len(last_table.rows) >= 1:
                header_cells = [c.text for c in last_table.rows[0].cells]
                fmt, idx_version, idx_description = classify_table_header(header_cells)
                signals["table_format"] = fmt
                if fmt == "new" and len(last_table.rows) >= 2:
                    # Every FILLED version in the table, in document order.
                    # Word tables routinely carry blank rows at the bottom, left
                    # over from a template or from someone tabbing past the end.
                    # Reading the literal last row then yielded "" for the
                    # current version — and because the comparison below drops
                    # empty sources before comparing, {header, "", sheet}
                    # collapsed to one value and the version-divergence check
                    # silently stopped running for that document. Found on a
                    # real base, where it had disabled the check without a
                    # single line of output saying so.
                    if idx_version is not None:
                        for row in last_table.rows[1:]:
                            cells = [c.text.strip() for c in row.cells]
                            if idx_version < len(cells) and cells[idx_version]:
                                signals["history_versions"].append(cells[idx_version])
                        if signals["history_versions"]:
                            signals["history_version"] = signals["history_versions"][-1]

                    # The description of that same last filled row.
                    if idx_description is not None:
                        for row in reversed(last_table.rows[1:]):
                            cells = [c.text.strip() for c in row.cells]
                            version_filled = (idx_version is None
                                              or (idx_version < len(cells) and cells[idx_version]))
                            if version_filled and idx_description < len(cells):
                                signals["history_description"] = cells[idx_description]
                                break
                elif fmt == "legacy" and len(last_table.rows) >= 2:
                    last_row = [c.text.strip() for c in last_table.rows[-1].cells]
                    if idx_description is not None and idx_description < len(last_row):
                        signals["history_description"] = last_row[idx_description]
                # format None: we don't try to extract anything — safer than
                # risking comparing the wrong column (e.g. mistaking Date for Version).

        # The footer's organization label, read here so the document is not
        # opened a second time just to get one line of it.
        try:
            footer_line = "".join(p.text for p in d.sections[0].footer.paragraphs)
        except (IndexError, AttributeError):
            footer_line = ""
        footer_label = footer_line.split("\t")[0].strip()
        if footer_label:
            signals["footer_label"] = footer_label

        hdr_paragraphs = [p.text for p in d.sections[0].header.paragraphs if p.text.strip()]
        if hdr_paragraphs:
            _parse_header(signals, hdr_paragraphs[0].strip(), full_path)
    except Exception as e:
        signals["error"] = str(e)
    return signals


def read_tab_category_map(wb):
    """
    Reads the Overview sheet's Category -> Tab mapping and returns it inverted
    (tab name -> real category name).

    This mapping is DATA, not something to recompute. The bootstrap may have
    had to alter a category name to make it a legal Excel sheet title (invalid
    characters, the 31-character cap, a collision), so deriving the tab name
    from the folder a second time here would mean keeping two implementations
    of the same transformation in sync forever — and when they drifted, the
    checker reported that a folder didn't match itself. Reading what the
    workbook already recorded removes that whole class of bug.

    Returns {} when there's no usable Overview, in which case the caller falls
    back to a direct comparison.
    """
    mapping = {}
    for sheet_name in wb.sheetnames:
        if not is_overview_sheet(sheet_name):
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            if not row or not row[0]:
                continue
            category = str(row[0]).strip()
            # Scan the row's cells from the right instead of trusting row[-1]:
            # openpyxl pads every row out to the sheet's widest column, so a
            # note typed into a column past "Tab" became the last cell, the
            # quoted tab name was never seen, and the mapping was lost — which
            # brought back exactly the false folder/tab mismatches on
            # sanitized or truncated tab names that this mapping exists to
            # prevent.
            for value in reversed(row[1:]):
                m = re.search(r"'([^']+)'", str(value or ""))
                if m:
                    mapping[m.group(1).strip()] = category
                    break
    return mapping


def strip_numeric_prefix_for_compare(folder_name):
    """'01 - Identity' -> 'Identity'. Used to compare a folder against its tab."""
    m = re.match(r"^\d+\s*-\s*(.+)$", (folder_name or "").strip())
    return m.group(1).strip() if m else (folder_name or "").strip()


def _style_name(paragraph):
    """
    The paragraph's style name, lowercased, or "" when it has none.

    python-docx returns None for `paragraph.style` whenever the paragraph
    points at a style id the document's styles part doesn't define. That is
    routine in real Word documents — ones converted from .doc, or with content
    pasted in from another file carrying its own styles. Reading `.name` off
    that None aborted the entire audit partway through, and because Python
    exits 1 on an uncaught exception, the crash was indistinguishable from this
    tool's own "problems found" exit code.

    Treating a style-less paragraph as having no style name is the conservative
    reading: it is then simply not recognized as a heading or a list, which is
    the same outcome as a paragraph styled Normal.
    """
    style = getattr(paragraph, "style", None)
    if style is None:
        return ""
    return (style.name or "").strip().lower()


def _style_chain(style, max_depth=10):
    """
    Yields a paragraph style and its base styles, outermost first. Word
    formatting cascades: a run may say nothing about bold/size/color while the
    paragraph's style (or the style that style is based on) does. Walking this
    chain is what keeps a visually correct document from being reported as a
    violation just because the property wasn't set on the run itself.
    """
    seen = 0
    while style is not None and seen < max_depth:
        yield style
        style = getattr(style, "base_style", None)
        seen += 1


def resolve_font_property(run, paragraph, prop):
    """
    Resolves an effective font property ("bold", "size", "color") for a run,
    following the cascade run -> paragraph style -> base styles.

    Returns (value, determined). `determined` is False when nothing in the
    cascade states the property explicitly — in that case the real value comes
    from the Word theme/defaults, which python-docx can't see reliably. In
    CONSERVATIVE mode we never call that a violation; it goes to the
    "could not verify" bucket instead.
    """
    def read(source_font):
        if prop == "bold":
            return source_font.bold
        if prop == "size":
            return source_font.size
        if prop == "color":
            color = source_font.color
            if color is None:
                return None
            try:
                return color.rgb
            except (AttributeError, ValueError):
                return None
        return None

    # run.bold is python-docx's own resolved view of the run's bold, so it is
    # preferred over reading the font directly for that one property.
    value = run.bold if prop == "bold" else read(run.font)
    if value is not None:
        return value, True

    for style in _style_chain(paragraph.style):
        font = getattr(style, "font", None)
        if font is None:
            continue
        value = read(font)
        if value is not None:
            return value, True

    return None, False


def _rgb_str(value):
    return str(value).upper() if value is not None else None


# A leading enumerator on a heading: "1. ", "2) ", "3 - ", "4: ". Stripped
# before the heading is matched against the section vocabulary, because that
# matching is startswith() against the bare term and could not see past it — a
# document whose sections read "1. Objetivo / 2. Pré-requisitos / 3. Passo a
# Passo" had every one of them reported missing. Numbering the sections is not
# something the standard forbids, and real bases do it.
_HEADING_ENUMERATOR = re.compile(r"^\s*\d+\s*[.)\-:]\s*")


def _heading_key(text):
    """A heading lowercased and stripped of a leading enumerator."""
    return _HEADING_ENUMERATOR.sub("", (text or "").strip().lower())


# Section 3 required headings, in the order the standard mandates. Optional
# sections (Roles & Responsibilities, Naming Convention, Verification) may sit
# between them, and an extra custom heading is allowed — the standard doesn't
# forbid one. What is checked: every required section exists, they appear in
# the mandated relative order, and Version History is last.
REQUIRED_SECTIONS = tuple(
    (key, _SECTIONS.get(key, ()))
    for key in ("purpose", "prerequisites", "steps", "history")
)


def _list_format(paragraph, doc):
    """
    Returns "bullet", "decimal", None (not a list), or "unknown" when the
    paragraph is a list item whose numbering definition can't be resolved.
    Word stores bullet and numbered lists under the same "List Paragraph"
    style, so the style name alone can't tell them apart — the numbering part
    has to be consulted.
    """
    # A paragraph can get its list formatting three ways. They are consulted in
    # WORD's order of precedence, which is not the order of convenience:
    #
    #   1. the paragraph's own w:numPr  — direct formatting, applied over all
    #   2. its style's w:numPr          — where List Bullet / List Number live
    #   3. the style's NAME             — last resort, when neither resolves
    #
    # Reading the style name first was wrong: a paragraph styled "List Number"
    # whose own numbering points at a bullet definition renders as a bullet,
    # and was reported as "decimal". That is exactly how the fixture's own
    # shared-numbering case managed to claim it modelled a numbered list while
    # the numbering part said bullet. The style name still has to stay as a
    # fallback, though — checking numbering alone flags documents using the
    # built-in list styles, whose numbering lives in the style definition, as
    # if they were not lists at all.
    sources = []
    pPr = paragraph._p.pPr
    direct = pPr.find(qn("w:numPr")) if pPr is not None else None
    if direct is not None:
        sources.append(direct)
    style_element = getattr(getattr(paragraph, "style", None), "element", None)
    style_pPr = style_element.find(qn("w:pPr")) if style_element is not None else None
    style_numPr = style_pPr.find(qn("w:numPr")) if style_pPr is not None else None
    if style_numPr is not None:
        sources.append(style_numPr)

    for numPr in sources:
        num_id_el = numPr.find(qn("w:numId"))
        if num_id_el is None:
            continue
        fmt = _numbering_format(doc, num_id_el.get(qn("w:val")))
        if fmt is not None:
            return "bullet" if fmt == "bullet" else "decimal"

    style_name = _style_name(paragraph)
    if "list bullet" in style_name:
        return "bullet"
    if "list number" in style_name:
        return "decimal"

    # A list item whose definition nothing could resolve is "unknown"; a
    # paragraph that is not a list item at all is None. The distinction is what
    # keeps an unresolvable list out of the violation count.
    return "unknown" if sources else None


def _has_page_field(footer):
    """
    True when the footer contains a Word PAGE field.

    Tested against the field instructions, not the raw XML. Searching the XML
    for the substring "PAGE" meant any footer whose TEXT happened to contain it
    — "ACME HOMEPAGE | Knowledge Base" — was credited with an automatic page
    field it doesn't have, so a typed-in page number went unreported. The word
    boundary matters too: " NUMPAGES " contains "PAGE", and a footer carrying
    only a total-pages field still has no current-page number.
    """
    element = getattr(footer, "_element", None)
    if element is None:
        return False
    for el in element.iter():
        if el.tag == qn("w:fldSimple"):
            instr = el.get(qn("w:instr")) or ""
        elif el.tag == qn("w:instrText"):
            instr = el.text or ""
        else:
            continue
        if re.search(r"\bPAGE\b", instr, re.IGNORECASE):
            return True
    return False


def _numbering_format(doc, num_id):
    """
    The resolved numFmt at level 0 for a numId ("decimal", "bullet", ...), or
    None when the numbering definition can't be resolved.
    """
    try:
        numbering = doc.part.numbering_part.element
    except (AttributeError, KeyError, ValueError):
        return None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) != num_id:
            continue
        abstract = num.find(qn("w:abstractNumId"))
        if abstract is None:
            return None
        abstract_id = abstract.get(qn("w:val"))
        for a in numbering.findall(qn("w:abstractNum")):
            if a.get(qn("w:abstractNumId")) != abstract_id:
                continue
            lvl = a.find(qn("w:lvl"))
            if lvl is None:
                return None
            fmt = lvl.find(qn("w:numFmt"))
            if fmt is None:
                return None
            return (fmt.get(qn("w:val")) or "").lower()
    return None


def _footer_signals(doc):
    """Returns (text, has_tab, has_page_field) for the first section's footer."""
    try:
        footer = doc.sections[0].footer
    except (IndexError, AttributeError):
        return None, False, False
    text = "".join(p.text for p in footer.paragraphs).strip()
    has_tab = any("\t" in p.text for p in footer.paragraphs)
    return text, has_tab, _has_page_field(footer)


def parse_version(text):
    """
    "1.10" -> (1, 10). Returns None when the value isn't a dotted number, so an
    unparseable cell is skipped rather than guessed at — the same conservative
    rule the formatting checks follow.

    Compared as integer tuples, not as text: "1.10" is a LATER version than
    "1.9", which a string comparison gets backwards.
    """
    if not text:
        return None
    cleaned = str(text).strip().lstrip("vV").strip()
    if not re.fullmatch(r"\d+(\.\d+)*", cleaned):
        return None
    return tuple(int(part) for part in cleaned.split("."))


def check_version_history_sequence(versions, label):
    """
    Checks Section 3's rules about the Version History as a SEQUENCE: the first
    row records the creation at 1.0, and versions only ever move forward.

    Both were previously unenforceable, because only the table's last row was
    ever read. A history missing its creation row, or one that goes backwards
    after a bad merge, is a document whose audit trail no longer reconstructs
    what happened to it — which is the entire purpose of the table.
    """
    issues = []
    if not versions:
        return issues  # no readable version column; reported elsewhere

    first = parse_version(versions[0])
    if first is None:
        return issues  # unparseable: don't guess
    if first != (1, 0):
        issues.append(
            "%s: the Version History starts at '%s'; Section 3 requires the first "
            "row to be the creation, at version 1.0. Either the creation row was "
            "deleted, or the table was started partway through the document's life."
            % (label, versions[0])
        )

    previous = first
    for current_text in versions[1:]:
        current = parse_version(current_text)
        if current is None:
            continue
        if current < previous:
            issues.append(
                "%s: the Version History goes backwards, from '%s' to '%s'. Rows "
                "are in chronological order, so each version must be greater than "
                "the one above it."
                % (label, ".".join(str(n) for n in previous), current_text)
            )
        previous = max(previous, current)
    return issues


def check_document_structure(d, label):
    """
    Checks Section 3: required sections present, in order, with Version History
    last. Takes an already-open Document; the caller reports an unreadable file
    once, on its own, rather than letting it surface here as a structure
    violation too.
    """
    issues = []
    headings = []
    for p in d.paragraphs:
        style_name = _style_name(p)
        if any(marker in style_name for marker in HEADING_STYLE_MARKERS):
            headings.append(p.text.strip())

    if not headings:
        return ["%s: no Heading 1 sections found — the document has no recognizable structure" % label]

    lowered = [_heading_key(h) for h in headings]
    positions = {}
    for key, synonyms in REQUIRED_SECTIONS:
        idx = next((i for i, h in enumerate(lowered)
                    if any(h.startswith(syn) for syn in synonyms)), None)
        if idx is None:
            issues.append(
                "%s: required section '%s' is missing (expected one of: %s)"
                % (label, key, ", ".join(synonyms))
            )
        else:
            positions[key] = idx

    ordered = [k for k, _ in REQUIRED_SECTIONS if k in positions]
    for earlier, later in zip(ordered, ordered[1:]):
        if positions[earlier] > positions[later]:
            issues.append(
                "%s: section '%s' appears after '%s'; the standard's order is "
                "Purpose -> Prerequisites -> Step by Step -> Version History"
                % (label, earlier, later)
            )

    if "history" in positions and positions["history"] != len(headings) - 1:
        trailing = headings[positions["history"] + 1:]
        issues.append(
            "%s: Version History must be the last section, but %s come after it"
            % (label, ", ".join("'%s'" % t for t in trailing))
        )

    return issues


def check_document_formatting(d, label):
    """
    Checks an already-open Document against the formatting constants above
    (Section 4 of the spec), in CONSERVATIVE mode: something is only reported as
    a violation when an explicit, divergent value was actually found. Anything
    the cascade can't settle is returned separately as "could not verify" and
    never counted as a problem — a checker that guesses produces false
    positives, and a formatting report nobody trusts is worse than no report at
    all.

    Returns (violations, unverifiable), both lists of strings.
    """
    violations = []
    unverifiable = []

    # --- margins (always determinable: they're stored on the section) ---
    for section in d.sections:
        for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            margin = getattr(section, side, None)
            if margin is None:
                unverifiable.append(f"{label}: '{side}' is not set explicitly")
                continue
            if abs(margin.inches - EXPECTED_MARGIN_INCHES) > 0.01:
                violations.append(
                    f"{label}: {side.replace('_', ' ')} is {margin.inches:.2f}in, "
                    f"expected {EXPECTED_MARGIN_INCHES}in"
                )
        break  # the standard is a single-section document; check the first

    # --- title: first non-empty paragraph of the body ---
    title_para = next((p for p in d.paragraphs if p.text.strip()), None)
    if title_para is None:
        unverifiable.append(f"{label}: no body paragraph found, title not checked")
    else:
        # Every run is checked, not just the first: a title split across runs
        # (autocorrect, pasted text) can have a correctly formatted first half
        # and an unformatted second half.
        for idx, run in enumerate(title_para.runs, 1):
            if not run.text.strip():
                continue
            where = f"{label}: title run {idx} ('{run.text.strip()[:30]}')"

            bold, determined = resolve_font_property(run, title_para, "bold")
            if not determined:
                unverifiable.append(f"{where}: bold not stated on the run or its style chain")
            elif bold is not True:
                violations.append(f"{where}: is not bold")

            size, determined = resolve_font_property(run, title_para, "size")
            if not determined:
                unverifiable.append(f"{where}: font size not stated on the run or its style chain")
            elif abs(size.pt - EXPECTED_TITLE_SIZE_PT) > 0.01:
                violations.append(
                    f"{where}: font size is {size.pt}pt, expected {EXPECTED_TITLE_SIZE_PT}pt"
                )

            color, determined = resolve_font_property(run, title_para, "color")
            if not determined or color is None:
                unverifiable.append(f"{where}: font color not stated on the run or its style chain")
            elif _rgb_str(color) != EXPECTED_TITLE_COLOR:
                violations.append(
                    f"{where}: font color is {_rgb_str(color)}, expected {EXPECTED_TITLE_COLOR}"
                )

    # --- section headings must use the NATIVE Heading 1 style ---
    for p in d.paragraphs:
        style_name = _style_name(p)
        if any(marker in style_name for marker in HEADING_STYLE_MARKERS):
            builtin = getattr(p.style, "builtin", None)
            if builtin is None:
                unverifiable.append(
                    f"{label}: heading '{p.text.strip()[:30]}' — couldn't tell whether the "
                    "style is native"
                )
            elif not builtin:
                violations.append(
                    f"{label}: heading '{p.text.strip()[:30]}' uses a custom style "
                    f"('{_style_name(p)}'), not Word's native Heading 1"
                )

    # --- warning callouts ---
    for p in d.paragraphs:
        if not p.text.strip().startswith(CALLOUT_PREFIX):
            continue
        snippet = p.text.strip()[:30]
        pPr = p._p.pPr
        shd = pPr.find(qn("w:shd")) if pPr is not None else None
        if shd is None:
            violations.append(
                f"{label}: callout '{snippet}' has no shading; expected "
                f"{EXPECTED_CALLOUT_SHADING}"
            )
        else:
            fill = (shd.get(qn("w:fill")) or "").upper()
            if fill != EXPECTED_CALLOUT_SHADING:
                violations.append(
                    f"{label}: callout '{snippet}' shading is {fill or 'empty'}, "
                    f"expected {EXPECTED_CALLOUT_SHADING}"
                )
        for idx, run in enumerate(p.runs, 1):
            if not run.text.strip():
                continue
            color, determined = resolve_font_property(run, p, "color")
            if not determined or color is None:
                unverifiable.append(
                    f"{label}: callout '{snippet}' run {idx} — text color not stated"
                )
            elif _rgb_str(color) != EXPECTED_CALLOUT_TEXT_COLOR:
                violations.append(
                    f"{label}: callout '{snippet}' run {idx} text color is "
                    f"{_rgb_str(color)}, expected {EXPECTED_CALLOUT_TEXT_COLOR}"
                )


    # --- page orientation ---
    try:
        from docx.enum.section import WD_ORIENT
        orientation = d.sections[0].orientation
        if orientation is not None and orientation != WD_ORIENT.PORTRAIT:
            violations.append("%s: page orientation is landscape, expected portrait" % label)
    except (ImportError, IndexError, AttributeError):
        unverifiable.append("%s: page orientation could not be read" % label)

    # --- footer: must exist, use a tab, and use Word's automatic page fields ---
    footer_text, footer_has_tab, footer_has_page = _footer_signals(d)
    if not footer_text:
        violations.append(
            "%s: the document has no footer; the standard requires "
            "'<Organization label>  |  Knowledge Base<TAB>Page <n> of <total>'" % label
        )
    else:
        if not footer_has_tab:
            violations.append(
                "%s: the footer has no tab separating the label from the page number" % label
            )
        if not footer_has_page:
            violations.append(
                "%s: the footer's page number is typed in rather than a Word automatic "
                "page field, so it won't update" % label
            )

    # --- header: plain hyphen between code and name, never an en/em dash ---
    try:
        header_text = "".join(p.text for p in d.sections[0].header.paragraphs)
    except (IndexError, AttributeError):
        header_text = ""
    # Only the separator BETWEEN THE CODE AND THE NAME is constrained. Scanning
    # the whole header flagged a document whose own name contains an em dash
    # ("KB-004 - Alterar Data de Expiração — Contingent Worker no
    # SuccessFactors"), which reads as an instruction to strip punctuation out
    # of the title — something Section 4 never asked for. A header with no code
    # has no such separator, so nothing is checked.
    separator = re.match(r"^\s*KB-\d+\s*([%s])" % "".join(DASHES), header_text)
    if separator:
        violations.append(
            "%s: the header uses '%s' between code and name; the standard "
            "requires a plain hyphen '-'" % (label, separator.group(1))
        )

    # --- list formatting per section, only where the section has content ---
    headings = []
    for i, p in enumerate(d.paragraphs):
        style_name = _style_name(p)
        if any(marker in style_name for marker in HEADING_STYLE_MARKERS):
            # The original text is what the report shows; the stripped key is
            # what gets matched, so "3. Passo a Passo" is recognized as the
            # Step by Step section instead of being skipped unchecked.
            headings.append((i, p.text.strip().lower(), _heading_key(p.text)))

    for pos, (idx, title, key) in enumerate(headings):
        accepted = None
        if any(key.startswith(x) for x in ANY_LIST_SECTIONS):
            accepted = ("bullet", "decimal")
        elif any(key.startswith(x) for x in NUMBERED_SECTIONS):
            accepted = ("decimal",)
        if accepted is None:
            continue
        end = headings[pos + 1][0] if pos + 1 < len(headings) else len(d.paragraphs)
        body_paras = [p for p in d.paragraphs[idx + 1:end] if p.text.strip()]
        if not body_paras:
            continue  # an empty section can't be judged
        formats = {_list_format(p, d) for p in body_paras}
        if "unknown" in formats and not formats & set(accepted):
            unverifiable.append(
                "%s: section '%s' — list numbering definition could not be resolved"
                % (label, title)
            )
        elif not formats & set(accepted):
            kind = ("a list (bulleted or numbered)" if len(accepted) > 1
                    else "a numbered list")
            violations.append("%s: section '%s' is not formatted as %s" % (label, title, kind))

    # --- each section gets its OWN numbering sequence ---
    # Word numbers every paragraph sharing a numId as one continuous list, in
    # document order. So when two sections share one, the second carries on from
    # the first ("Step by Step - ServiceNow" starting at 9) instead of restarting
    # at 1. Nothing about the document looks wrong while you read it: Word shows
    # the numbers, they are simply the wrong ones, and the defect only surfaces
    # when someone follows step 9 of a procedure that has four steps.
    per_section = []
    for p in d.paragraphs:
        if any(m in _style_name(p) for m in HEADING_STYLE_MARKERS):
            per_section.append((p.text.strip(), set()))
            continue
        if not per_section:
            continue
        pPr = p._p.pPr
        numPr = pPr.find(qn("w:numPr")) if pPr is not None else None
        if numPr is None:
            continue
        num_id = numPr.find(qn("w:numId"))
        if num_id is not None and num_id.get(qn("w:val")):
            value = num_id.get(qn("w:val"))
            # Only ORDERED lists can carry a count from one section into the
            # next. Bullets share a numbering definition all the time — Word
            # reuses one for identical bullet formatting — and sharing it is
            # harmless, because there is no number on screen to be wrong. On a
            # real base this was 9 of 19 findings from this check, every one of
            # them false. An unresolvable definition is left alone too, the
            # same conservative rule the rest of the formatting checks follow.
            fmt = _numbering_format(d, value)
            if fmt is not None and fmt not in ("bullet", "none"):
                per_section[-1][1].add(value)

    owner = {}
    for heading, num_ids in per_section:
        for num_id in sorted(num_ids):
            first = owner.setdefault(num_id, heading)
            if first != heading:
                violations.append(
                    "%s: sections '%s' and '%s' share one numbering sequence, so the "
                    "second continues the first's count instead of restarting at 1"
                    % (label, first[:34], heading[:34])
                )

    # --- author names in the Version History must be Title Case, not all caps ---
    if d.tables:
        last_table = d.tables[-1]
        if len(last_table.rows) >= 2:
            header_cells = [c.text for c in last_table.rows[0].cells]
            fmt, _idx_v, _idx_d = classify_table_header(header_cells)
            if fmt is not None:
                lowered = [c.strip().lower() for c in header_cells]
                author_idx = next(
                    (i for i, c in enumerate(lowered)
                     if "autor" in c or "author" in c or "revisor" in c or "reviewer" in c),
                    None,
                )
                if author_idx is not None:
                    for row in last_table.rows[1:]:
                        if author_idx >= len(row.cells):
                            continue
                        value = row.cells[author_idx].text.strip()
                        letters = [ch for ch in value if ch.isalpha()]
                        if len(letters) > 1 and value == value.upper():
                            violations.append(
                                "%s: Version History author '%s' is in all caps; the "
                                "standard requires Title Case" % (label, value)
                            )

    return violations, unverifiable


def resolve_missing_master_list(base_dir):
    """
    Called when no master-index spreadsheet was found. Interactively offers to
    create one (delegating to bootstrap_master_list.py) or to be pointed at an
    existing file.

    IMPORTANT: this only prompts when stdin is an interactive terminal. Under
    CI, cron, or a pipe there is nobody to answer, and a script that blocks
    waiting for input would hang the pipeline — in that case it prints the same
    guidance and the caller exits with code 2. This is also why creation lives
    in a separate script: check_master_list.py itself never writes.

    Returns (path, created) — created is True only when a spreadsheet was
    actually generated in this run. Pointing at an existing file must not
    claim anything was written. (None, False) means abort.
    """
    print(f"No master-index spreadsheet found in '{base_dir}'.")
    print(f"(looked for: {', '.join(MASTER_LIST_NAME_PATTERNS)})\n")

    def print_guidance():
        print(
            "Running non-interactively, so I won't prompt. Either:\n\n"
            "  1) generate one from the files already in the base:\n"
            f"       python3 bootstrap_master_list.py \"{base_dir}\"\n\n"
            "  2) or point at an existing spreadsheet:\n"
            f"       python3 check_master_list.py \"{base_dir}\" \"<file name>.xlsx\"\n"
        )

    # isatty() is necessary but not sufficient: on Windows, stdin redirected
    # from NUL is reported as a TERMINAL, so a cron/CI/pipeline run took the
    # interactive path below and died on its first prompt instead of printing
    # this guidance. An immediate EOF is therefore treated the same way.
    if not stdin_is_interactive():
        print_guidance()
        return None, False

    print("What would you like to do?")
    print("  [1] Create one now by scanning this base")
    print("  [2] Point me at an existing spreadsheet")
    print("  [3] Abort")

    try:
        choice = input("> ").strip()
    except EOFError:
        print_guidance()
        return None, False
    except KeyboardInterrupt:
        print("\nAborted.")
        return None, False

    if choice == "1":
        try:
            from bootstrap_master_list import bootstrap, HEADERS, STATUS_IN_REVIEW
        except ImportError as e:
            print(f"Couldn't load bootstrap_master_list.py ({e}).")
            return None, False

        # Ask which language the base is authored in. Defaulting silently to
        # English would write an index whose column headers and status values
        # don't match the documents it indexes.
        #
        # Offer only what actually exists: the old prompt invited "another
        # language — give its two-letter code", then died on HEADERS[lang] for
        # anything but en/pt and surfaced it as a bare "ERROR: 'fr'". The CLI's
        # --lang has always validated; this path did not.
        supported = sorted(HEADERS)
        print("\nWhat language are this base's documents written in?")
        for i, code in enumerate(supported, 1):
            print("  [%d] %s%s" % (i, code, "  (default)" if code == "en" else ""))
        try:
            answer = (input("> ").strip().lower() or "en")
        except (EOFError, KeyboardInterrupt):
            print("\nAborted.")
            return None, False
        if answer.isdigit() and 1 <= int(answer) <= len(supported):
            lang = supported[int(answer) - 1]
        elif answer in HEADERS:
            lang = answer
        else:
            print(
                "Unsupported language %r. Supported: %s. To add one, extend the "
                "HEADERS/OVERVIEW_HEADERS/STATUS_* tables in "
                "bootstrap_master_list.py. Nothing was written."
                % (answer, ", ".join(supported))
            )
            return None, False

        # The label shown here has to be the one that will actually be written:
        # this message used to name the Portuguese status regardless of the
        # language chosen.
        print(
            "\nHeads-up: every row will be written with Status '%s', and any\n"
            "document whose file name has no KB code will be left with '%s' instead of\n"
            "a new code — codes are never recycled, so that choice stays with you.\n"
            % (STATUS_IN_REVIEW[lang], NO_CODE_MARKER)
        )
        try:
            confirm = input("Proceed? [y/N] ").strip().lower()
        except (EOFError, KeyboardInterrupt):
            print("\nAborted.")
            return None, False
        if confirm not in ("y", "yes", "s", "sim"):
            print("Aborted. Nothing was written.")
            return None, False
        try:
            return bootstrap(base_dir, lang=lang), True
        except Exception as e:
            print(f"ERROR: {e}")
            return None, False

    if choice == "2":
        try:
            name = input("File name (inside the base folder): ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\nAborted.")
            return None, False
        candidate = os.path.join(base_dir, name)
        if not os.path.isfile(candidate):
            print(f"ERROR: '{name}' not found in '{base_dir}'.")
            return None, False
        return candidate, False

    print("Aborted. Nothing was written.")
    return None, False


def find_non_ascii_filenames(by_path):
    """
    Lists (informational, NEVER an error) the file paths that contain a
    non-ASCII character (accents, etc.).

    This is different from the NFC/NFD bug handled in nfc(): here the
    Unicode form doesn't matter, only whether the name has any character
    outside the basic ASCII range. It's a PORTABILITY note (the name could
    break if the base is ever exported to a system that doesn't handle
    Unicode well), not a base inconsistency — so it never enters the
    problem count, it's just a heads-up for a possible future migration.
    """
    found = []
    for path in sorted(by_path.keys()):
        if not path.isascii():
            found.append(path)
    return found


def format_date(dt):
    if dt is None:
        return "unknown"
    if isinstance(dt, datetime.datetime):
        return dt.strftime("%Y-%m-%d")
    return str(dt)


def check(base_dir, master_list_name=None):
    bootstrapped = False
    try:
        master_path = find_master_list(base_dir, master_list_name)
    except FileNotFoundError:
        # No spreadsheet: offer to create one or be pointed at it, instead of
        # dying with a traceback. Returns None to signal "couldn't run".
        master_path, bootstrapped = resolve_missing_master_list(base_dir)
        if master_path is None:
            return None
    print(f"Base: {base_dir}")
    print(f"Master-index spreadsheet: {os.path.basename(master_path)}\n")

    by_basename, by_path = index_real_files(base_dir)

    # Which way round a typed "03/04/2026" should be read. Recorded by the
    # bootstrap in .kb-compiler.json; day-first stays the default, since that is
    # what every base written before this existed was already read as.
    base_language = str(read_config(base_dir).get("language") or "").strip().lower()
    day_first = not base_language.startswith("en")

    wb = openpyxl.load_workbook(master_path, data_only=True)
    tab_to_category = read_tab_category_map(wb)

    broken_reference_issues = []
    duplicate_code_issues = []
    name_divergence_issues = []
    version_divergence_issues = []
    legacy_table_issues = []
    header_missing_tab_issues = []
    header_missing_version_issues = []
    invalid_status_issues = []
    missing_code_issues = []
    version_sequence_issues = []
    unreadable_document_issues = []
    formatting_issues = []
    formatting_unverifiable = []
    structure_issues = []
    folder_mismatch_issues = []
    ambiguous_reference_issues = []
    footer_labels = {}
    codes_seen = {}
    referenced_files = set()
    sheets_without_header = []

    for sheet_name in wb.sheetnames:
        if is_overview_sheet(sheet_name):
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        # Find the column-header row by name instead of assuming it's always
        # row 2 — a sheet that doesn't have a recognizable "Code"/"Código"
        # column is skipped (with a warning) rather than silently reading
        # the wrong columns as data.
        header_idx, header_cells = find_header_row(rows)
        if header_idx is None:
            sheets_without_header.append(sheet_name)
            continue
        col = map_columns(header_cells)
        if "code" not in col or "document" not in col or "file" not in col:
            sheets_without_header.append(sheet_name)
            continue

        def cell(row, key):
            idx = col.get(key)
            return row[idx] if idx is not None and idx < len(row) else None

        for row in rows[header_idx + 1:]:
            if not row:
                continue
            raw_code = cell(row, "code")
            codigo = str(raw_code).strip() if raw_code is not None else ""
            if not codigo:
                # A row with real content but an empty Code used to be dropped
                # in silence: nothing reported the row, and the document it
                # indexes then resurfaced as an "orphan", pointing at the file
                # instead of at the row that is actually wrong. Section 2 has
                # the "—" marker precisely so a document without a code is
                # still visible in the index. A whitespace-only cell counts as
                # empty here, since it reads as blank to everyone.
                if cell(row, "document") or cell(row, "file"):
                    missing_code_issues.append(
                        "[%s] the row for '%s' (file '%s') has an empty Code "
                        "cell, so the whole row was skipped. Section 2 requires "
                        "a code, or the '%s' marker for a document that "
                        "legitimately has none — never a blank."
                        % (sheet_name, cell(row, "document"),
                           cell(row, "file"), NO_CODE_MARKER)
                    )
                # The row is broken, but it DOES point at this file, so the
                # file is not unindexed. Recording it keeps one defect to one
                # finding, instead of also reporting the document as an orphan
                # and sending the reader after the wrong thing — the same rule
                # the unreadable-document handling below follows.
                referenced = cell(row, "file")
                if referenced:
                    referenced = nfc(str(referenced))
                    if "/" in referenced:
                        referenced_files.add(referenced)
                    else:
                        for match in by_basename.get(referenced, []):
                            referenced_files.add(match)
                continue
            documento = cell(row, "document")
            arquivo = cell(row, "file")
            status = cell(row, "status")
            versao_planilha = cell(row, "version")
            ultima_atualizacao = cell(row, "last_updated")

            # 1) duplicate code
            if codigo != NO_CODE_MARKER:
                if codigo in codes_seen:
                    duplicate_code_issues.append(
                        f"Code '{codigo}' appears in '{codes_seen[codigo]}' and in '{sheet_name}' "
                        f"(current document: '{documento}')"
                    )
                else:
                    codes_seen[codigo] = sheet_name

            # 1b) Status must be one of the three values Section 7 allows. This
            # is checked on the ROW, before the file is even resolved, because
            # Status is a property of the index rather than of the document.
            status_text = str(status).strip() if status is not None else ""
            if not status_text:
                invalid_status_issues.append(
                    "[%s] %s ('%s'): the Status cell is empty. Section 7 requires "
                    "one of: %s." % (sheet_name, codigo, documento,
                                     ", ".join(STATUS_WRITTEN))
                )
            elif status_text.lower() not in VALID_STATUS_VALUES:
                invalid_status_issues.append(
                    "[%s] %s ('%s'): Status is '%s', which is not one of the values "
                    "Section 7 allows (%s). A value the tooling doesn't recognize "
                    "silently breaks things rather than failing loudly: the "
                    "Overview's COUNTIF formulas stop counting this row, and a "
                    "misspelled legacy label makes the document fall through to "
                    "the structure checks it should be exempt from."
                    % (sheet_name, codigo, documento, status_text,
                       ", ".join(STATUS_WRITTEN))
                )

            if not arquivo:
                continue
            # normalize to NFC before any file-name comparison, so "same name,
            # different Unicode form" isn't mistaken for a real broken
            # reference (see nfc()).
            arquivo = nfc(str(arquivo))

            # 2) broken reference
            full_path = None
            resolved_rel = None
            if "/" in arquivo:
                full_path = by_path.get(arquivo)
                resolved_rel = arquivo if full_path else None
            else:
                matches = by_basename.get(arquivo) or []
                if len(matches) == 1:
                    resolved_rel = matches[0]
                    full_path = by_path.get(resolved_rel)
                elif len(matches) > 1:
                    # The same file name exists in more than one folder. Taking
                    # the first match would silently compare the WRONG document
                    # — its title, header and version would come from a file
                    # this row doesn't refer to. The row's tab tells us which
                    # category it belongs to, so use that to narrow it down.
                    expected_category = tab_to_category.get(sheet_name, sheet_name)
                    expected_label = strip_numeric_prefix_for_compare(expected_category).lower()
                    narrowed = [
                        m for m in matches
                        if strip_numeric_prefix_for_compare(m.split("/")[0]).lower()
                        == expected_label
                    ]
                    if len(narrowed) == 1:
                        resolved_rel = narrowed[0]
                        full_path = by_path.get(resolved_rel)
                    else:
                        ambiguous_reference_issues.append(
                            "[%s] %s ('%s'): the File column says '%s', but that name exists "
                            "in %d folders (%s) and the row's tab doesn't single one out. "
                            "Nothing about this document was checked. Use the full relative "
                            "path in the File column, or rename one of the files."
                            % (sheet_name, codigo, documento, arquivo, len(matches),
                               ", ".join("'%s'" % m for m in matches))
                        )
                        for m in matches:
                            referenced_files.add(m)
                        continue

            if resolved_rel:
                # Record the resolved path, not the bare name: otherwise a
                # same-named file in another folder that nothing indexes would
                # escape orphan detection by borrowing this row's reference.
                referenced_files.add(resolved_rel)

            if not full_path:
                broken_reference_issues.append(
                    f"[{sheet_name}] {codigo} ('{documento}'): master-index spreadsheet points to "
                    f"'{arquivo}', which doesn't exist in the base."
                )
                continue

            if not full_path.lower().endswith(".docx"):
                continue  # only .docx has a title/header/history to compare

            label = f"[{sheet_name}] {codigo}"

            # Open the document ONCE and share it with every check. It used to
            # be parsed four times per row (signals, structure, formatting, and
            # again for the footer label). Opening it up front also means an
            # unreadable file is reported once, under its own heading: before,
            # a single corrupt .docx produced THREE findings — a broken
            # reference (claiming a file that plainly exists doesn't), a
            # structure violation and a formatting violation.
            try:
                document = docx.Document(full_path)
            except Exception as e:
                unreadable_document_issues.append(
                    f"{label} ('{documento}'): '{arquivo}' exists but couldn't be "
                    f"opened as a Word document ({e}). Nothing about it was checked."
                )
                continue

            signals = read_docx_signals(full_path, doc=document)
            if signals["error"]:
                unreadable_document_issues.append(
                    f"{label} ('{documento}'): '{arquivo}' opened, but its contents "
                    f"couldn't be read ({signals['error']}). Nothing about it was checked."
                )
                continue

            fmt_violations, fmt_unverifiable = check_document_formatting(document, label)
            formatting_issues.extend(fmt_violations)
            formatting_unverifiable.extend(fmt_unverifiable)
            # Legacy documents are, by definition, the ones that predate the
            # standard — flagging every one of them for missing sections would
            # flood the report with findings nobody intends to fix, which is
            # exactly the failure the conservative mode exists to avoid.
            if str(status or "").strip().lower() not in LEGACY_STATUS_LABELS:
                structure_issues.extend(check_document_structure(document, label))
                # Same exemption as the structure checks, and for the same
                # reason: a superseded document predates the standard, so its
                # history was never going to start at 1.0.
                version_sequence_issues.extend(
                    check_version_history_sequence(signals["history_versions"], label)
                )

            # The footer's organization label. It can't be validated in
            # isolation (the script doesn't know the base's label), but every
            # document in one base must carry the SAME one.
            if signals["footer_label"]:
                footer_labels.setdefault(signals["footer_label"], []).append(codigo)

            # Folder/tab mismatch. The File column stores only the bare name
            # precisely because the tab is supposed to imply the folder; if the
            # file actually sits somewhere else, that assumption is broken and
            # the row is wrong, even though nothing else complains.
            rel = os.path.relpath(full_path, base_dir).replace(os.sep, "/")
            folder = rel.split("/")[0] if "/" in rel else ""
            folder_label = strip_numeric_prefix_for_compare(folder)
            # The tab may be a sanitized/truncated version of the real category
            # name; the Overview records which category each tab represents.
            expected = tab_to_category.get(sheet_name, sheet_name)
            expected_label = strip_numeric_prefix_for_compare(expected)
            # A superseded document in the Legacy folder whose row stayed on its
            # topical tab is EXPLICITLY allowed by Section 7 ("they can stay
            # listed in their original topical category tab (with Status =
            # 'Legacy') or get their own 'Legacy' tab"). Reporting it told the
            # reader to undo an arrangement the standard offers them, which is
            # worse than saying nothing at all.
            is_legacy_row = str(status or "").strip().lower() in LEGACY_STATUS_LABELS
            in_legacy_folder = folder_label.strip().lower() in LEGACY_FOLDER_NAMES
            if (folder_label
                    and folder_label.lower() != expected_label.lower()
                    and not (is_legacy_row and in_legacy_folder)):
                folder_mismatch_issues.append(
                    f"{label} ('{documento}'): the row is filed under category "
                    f"'{expected_label}', but the file is in folder '{folder}'. "
                    "Either move the file or move the row."
                )

            # 3a) Version History table in the legacy format (no real semantic
            # version column) — reported as its own item, instead of producing
            # a false "version divergence" comparing something the table doesn't have.
            if signals["table_format"] == "legacy":
                legacy_table_issues.append(
                    f"[{sheet_name}] {codigo} ('{documento}'): the Version History table is "
                    "in the old format (No. | Revision Date | Revision | Reviewer), which has no "
                    "semantic-version column — just a revision counter. "
                    "Action needed: convert the table to the new format "
                    "(Author | Version | Date | Change Description)."
                )

            # 3b) header missing the tab separator between name and version
            # (e.g. 'Namev1.2' glued together instead of 'Name<TAB>v1.2').
            # Reported as its own item so it doesn't show up as a confusing
            # 'header: None' inside the generic version divergence.
            if signals["header_missing_tab"]:
                header_missing_tab_issues.append(
                    f"[{sheet_name}] {codigo} ('{documento}'): the header has no tab separator "
                    f"between name and version (found: '{signals['header']}'). "
                    "Action needed: fix the header to the format "
                    "'KB-XXX - <Document Name><TAB>v<Version>', with a real tab character "
                    "separating the name from the version."
                )

            # 3c) header that carries no version at all. Section 4 requires one
            # in every header, but an ABSENT version used to be invisible: the
            # version comparison below drops empty sources before comparing, so
            # {None, "1.0", "1.0"} collapsed to a single value and reported
            # nothing. A header with no version is exactly the case where the
            # "header must match the Version History" rule silently stops being
            # enforceable.
            if signals["header"] is None:
                header_missing_version_issues.append(
                    f"{label} ('{documento}'): the document has no header at all. "
                    "Section 4 requires 'KB-XXX - <Document Name><TAB>v<Version>'."
                )
            elif signals["header_missing_version"]:
                header_missing_version_issues.append(
                    f"{label} ('{documento}'): the header carries no version "
                    f"(found: '{signals['header']}'). Section 4 requires "
                    "'KB-XXX - <Document Name><TAB>v<Version>', and that version must "
                    "match the last row of the Version History."
                )

            file_name = strip_code_prefix(os.path.splitext(os.path.basename(arquivo))[0])
            title_name = strip_code_prefix(signals["title"])
            header_name = strip_code_prefix(signals["header_name"])
            spreadsheet_name = str(documento).strip() if documento else None

            candidates = {
                "file name": file_name,
                "title (document body)": title_name,
                "header": header_name,
                "master-index spreadsheet (Document column)": spreadsheet_name,
            }
            # Compared NFC-normalized, for the same reason every file-name
            # comparison is (see nfc()): the file side was already normalized,
            # but the title, header and spreadsheet sides were not — so an
            # accented name authored as NFD and recorded as NFC was reported as
            # a name divergence between two strings that are the same text,
            # with a recommendation to "fix" one of them.
            unique_names = set(nfc(v) for v in candidates.values() if v)
            if len(unique_names) > 1:
                mtime = os.path.getmtime(full_path)
                mtime_dt = datetime.datetime.fromtimestamp(mtime)
                spreadsheet_date = parse_date_value(ultima_atualizacao,
                                                    day_first=day_first)
                ambiguity = ""
                if date_is_ambiguous(ultima_atualizacao):
                    # Disclosed rather than resolved in silence: the reading
                    # chosen here is what decides which source the
                    # recommendation points at, and the other reading can point
                    # at the other source.
                    ambiguity = (
                        " (note: '%s' is ambiguous — read as %s under this "
                        "base's %s-first convention; the other reading is a "
                        "different date, so confirm before relying on this)"
                        % (ultima_atualizacao, format_date(spreadsheet_date),
                           "day" if day_first else "month"))

                if ultima_atualizacao is not None and spreadsheet_date is None:
                    recommendation = (
                        "no automatic recommendation — the master-index spreadsheet's "
                        f"Last Updated value ('{ultima_atualizacao}') couldn't be read as a date"
                    )
                elif spreadsheet_date is None:
                    recommendation = "no automatic recommendation — manual decision needed"
                elif mtime_dt.date() > spreadsheet_date:
                    recommendation = (
                        f"use the DOCUMENT's name (title/header), since the file was "
                        f"last modified on {format_date(mtime_dt)}, after the last "
                        f"update recorded in the master-index spreadsheet ({format_date(spreadsheet_date)})"
                    )
                elif spreadsheet_date > mtime_dt.date():
                    recommendation = (
                        f"use the MASTER-INDEX SPREADSHEET's name, since it was updated on "
                        f"{format_date(spreadsheet_date)}, after the file's last "
                        f"modification ({format_date(mtime_dt)})"
                    )
                else:
                    recommendation = "no automatic recommendation — manual decision needed"

                recommendation += ambiguity
                detail = [f"[{sheet_name}] {codigo}: names diverge between sources:"]
                for origin, value in candidates.items():
                    detail.append(f"    - {origin}: '{value}'")
                detail.append(f"    Recommendation: {recommendation}")
                detail.append(
                    "    Action needed: pick which of these variants becomes the official name and "
                    "apply it in all 4 sources (file, title, header, master-index spreadsheet)."
                )
                detail.append(
                    "    Note: applying that choice to all 4 sources is a content edit, not just a "
                    "name change — it triggers the critical version rule (the header must always "
                    "match the last row of the Version History). In other words, resolving this name "
                    "divergence also requires bumping the document's version and adding a new row to "
                    "the Version History describing the name standardization — otherwise fixing one "
                    "inconsistency creates another."
                )
                name_divergence_issues.append("\n".join(detail))

            # 3) version divergence
            versions = {
                "header": signals["header_version"],
                "last row of Version History": signals["history_version"],
                "master-index spreadsheet (Version column)": str(versao_planilha).strip() if versao_planilha else None,
            }
            unique_versions = set(v for v in versions.values() if v)
            if len(unique_versions) > 1:
                detail = [f"[{sheet_name}] {codigo}: versions diverge between sources:"]
                for origin, value in versions.items():
                    detail.append(f"    - {origin}: '{value}'")
                detail.append(
                    "    Action needed: align all 3 sources to the most recent version "
                    "(usually the Version History's, since it's the audit record)."
                )
                version_divergence_issues.append("\n".join(detail))

    # 4) orphaned files (exist in the base, but aren't in any master-index row)
    orphan_file_issues = []
    for basename, paths in by_basename.items():
        # The spreadsheet in use is never an orphan — including when it was
        # passed by an explicit name that doesn't match the usual patterns.
        if (is_master_list_filename(basename)
                or basename == nfc(os.path.basename(master_path))
                or basename == CONFIG_FILENAME):
            continue
        for path in paths:
            if is_inside_attachment_folder(path):
                continue  # attachment of a KB document, not a standalone indexable item
            # Resolved paths only. The old `or basename not in referenced_files`
            # fallback could never fire (a bare name only enters the set for a
            # root-level file, which by then has no unindexed twin) and
            # re-opened the very hole 0.8.1 closed.
            if path not in referenced_files:
                orphan_file_issues.append(f"'{path}' exists in the base, but isn't in any master-index row")

    # A skipped sheet is a finding, not a footnote. It used to print only a
    # WARNING and contribute nothing to the count, so a base whose sheet had
    # lost its header row reported "SUMMARY: 0 problem(s) found" and exited 0 —
    # the documented "clean" signal for a CI gate — while an entire sheet went
    # unchecked. Whether the exit code happened to be non-zero depended on
    # orphan detection incidentally firing for the same documents.
    skipped_sheet_issues = [
        "sheet '%s' has no recognizable column-header row (a 'Code'/'Código' "
        "column), so it was skipped entirely and nothing it indexes was "
        "checked. This run cannot vouch for those documents." % name
        for name in sheets_without_header
    ]

    def section(title, items):
        print(f"=== {title} ({len(items)}) ===")
        if not items:
            print("No problems found.\n")
            return
        for item in items:
            print("- " + item)
        print()

    section("Sheets skipped entirely (no recognizable column-header row)",
            skipped_sheet_issues)
    section("Broken references (master-index spreadsheet → file)", broken_reference_issues)
    section("Orphaned files (file → master-index spreadsheet)", orphan_file_issues)
    section("Duplicate/shared codes", duplicate_code_issues)
    section("Name divergence (file/title/header/master-index spreadsheet)", name_divergence_issues)
    section("Version divergence (header/history/master-index spreadsheet)", version_divergence_issues)
    section("Version History table in the legacy format", legacy_table_issues)
    section("Header missing the tab separator (name and version glued or space-separated)", header_missing_tab_issues)
    section("Header missing its version (Section 4)", header_missing_version_issues)
    section("Unreadable documents (the file exists but cannot be parsed)",
            unreadable_document_issues)
    section("Invalid Status value (Section 7)", invalid_status_issues)
    section("Row with an empty Code cell (Section 2)", missing_code_issues)
    section("Version History sequence (Section 3)", version_sequence_issues)
    if len(footer_labels) > 1:
        detail = ["the base uses more than one footer label; every document must carry the same one:"]
        for prefix, codes in sorted(footer_labels.items()):
            shown = ", ".join(codes[:4]) + (", ..." if len(codes) > 4 else "")
            detail.append("    - '%s'  (%s)" % (prefix, shown))
        formatting_issues.append("\n".join(detail))

    section("Structure violations (Section 3 of the standard)", structure_issues)
    section("Ambiguous file reference (same name in several folders)", ambiguous_reference_issues)
    section("Folder/tab mismatch", folder_mismatch_issues)
    section("Formatting violations (Section 4 of the standard)", formatting_issues)

    # Informational — NOT a problem, and not part of the SUMMARY. Just a map
    # of files with a non-ASCII character in the name, in case of a future
    # migration to a system that doesn't handle Unicode well.
    # Informational — NOT counted as problems. Conservative mode: when the
    # run/style cascade doesn't state a formatting property explicitly, the
    # real value comes from Word's theme defaults, which can't be read
    # reliably here. Reporting those as violations would flood the report
    # with false positives, so they're listed apart instead.
    print(f"--- Informational: formatting properties that could not be verified "
          f"({len(formatting_unverifiable)}) ---")
    print(
        "These are NOT violations. The property simply isn't stated on the run or anywhere in\n"
        "its style chain, so its effective value comes from Word's defaults and can't be\n"
        "confirmed programmatically. Check visually if a specific document looks off."
    )
    for item in formatting_unverifiable:
        print("- " + item)
    print()

    non_ascii = find_non_ascii_filenames(by_path)
    print(f"--- Informational: files with a non-ASCII character in the name ({len(non_ascii)}) ---")
    print(
        "This is NOT a base consistency problem — it's only a portability note.\n"
        "Accented characters in file names work normally here; they'd only break if the base were\n"
        "ever exported to a system that doesn't handle Unicode well. Nothing to fix now."
    )
    if non_ascii:
        for path in non_ascii:
            print("- " + path)
    print()

    total = (
        len(broken_reference_issues)
        + len(orphan_file_issues)
        + len(duplicate_code_issues)
        + len(name_divergence_issues)
        + len(version_divergence_issues)
        + len(legacy_table_issues)
        + len(header_missing_tab_issues)
        + len(header_missing_version_issues)
        + len(unreadable_document_issues)
        + len(invalid_status_issues)
        + len(missing_code_issues)
        + len(skipped_sheet_issues)
        + len(version_sequence_issues)
        + len(formatting_issues)
        + len(structure_issues)
        + len(folder_mismatch_issues)
        + len(ambiguous_reference_issues)
    )
    if bootstrapped:
        # The checker itself still wrote nothing; the spreadsheet was created by
        # bootstrap_master_list.py, which this run delegated to after asking.
        # Saying "no file was changed" flat out would be false here.
        print(
            f"SUMMARY: {total} problem(s) found. A master-index spreadsheet was created "
            "for this base (see above); the checks themselves changed no file."
        )
    else:
        print(f"SUMMARY: {total} problem(s) found. No file was changed by this script.")
    return total


if __name__ == "__main__":
    if len(sys.argv) < 2:
        # Usage error is 'couldn't run' (2), not 'problems found' (1).
        print(__doc__)
        sys.exit(2)
    base_dir_arg = sys.argv[1]
    master_list_arg = sys.argv[2] if len(sys.argv) > 2 else None
    try:
        total_problems = check(base_dir_arg, master_list_arg)
    except (FileNotFoundError, RuntimeError) as e:
        print(f"ERROR: {e}")
        sys.exit(2)
    except Exception:
        # Any unexpected failure is "couldn't run" (2), never "problems found"
        # (1). Python exits 1 on an uncaught exception, and this tool's own
        # contract reads 1 as a completed audit that found things — so without
        # this, a crash silently passes for a result in any wrapper, CI job or
        # batch-edit gate that keys on the exit code. Found by running against
        # a real base: one document with a style-less paragraph took down the
        # whole audit and reported it as a finding.
        import traceback
        print("ERROR: the audit failed to complete. Nothing was changed.")
        print("Please report the traceback below.")
        traceback.print_exc()
        sys.exit(2)

    # Exit codes are distinct on purpose so a wrapper can tell the three cases
    # apart: 0 = clean, 1 = problems found, 2 = couldn't run at all.
    if total_problems is None:
        sys.exit(2)
    sys.exit(1 if total_problems else 0)
